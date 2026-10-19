"""Excel tiers report (openpyxl), the format the internship team worked from.

Sheets: Summary, High / Medium / Low priority, All Scored (with per-factor
contributions), Ineligible (with reasons), Duplicates (with the canonical tender
and Jaccard similarity).
"""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from bidalert.models import ScreenedTender
from bidalert.processing.deadlines import IST, days_remaining
from bidalert.profile import FACTORS, CompanyProfile

HEADER_FILL = PatternFill("solid", fgColor="1C5CAB")
HEADER_FONT = Font(bold=True, color="FFFFFF")

BASE_COLS = [
    "Tender ID",
    "Portal",
    "Title",
    "Organisation",
    "Category",
    "State",
    "Value (INR)",
    "Closing (IST)",
    "Days left",
    "Score",
    "Tier",
    "Unverified",
    "Link",
]


def _base(s: ScreenedTender, now: Any) -> list[Any]:
    t = s.tender
    return [
        t.source_tender_id,
        t.portal,
        t.title,
        t.organisation,
        s.category or "",
        t.state or "",
        t.estimated_value_inr,
        t.closing_at.astimezone(IST).replace(tzinfo=None),
        round(days_remaining(t.closing_at, now), 1),
        s.score.total if s.score else None,
        s.score.tier if s.score else "",
        ", ".join(s.eligibility.unverified) if s.eligibility else "",
        t.url or "",
    ]


def _sheet(
    wb: Workbook, title: str, header: Sequence[str], rows: Sequence[Sequence[Any]]
) -> Worksheet:
    ws = wb.create_sheet(title)
    ws.append(list(header))
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(list(r))
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, name in enumerate(header, start=1):
        width = (
            60 if name in ("Title", "Reasons") else 34 if name in ("Organisation", "Link") else 16
        )
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if header[cell.column - 1].startswith("Closing"):
                cell.number_format = "dd-mmm-yyyy hh:mm"
            elif header[cell.column - 1] == "Value (INR)":
                cell.number_format = "#,##0"
    return ws


def write_excel(result: Any, profile: CompanyProfile, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    now = result.now
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    data_note = "LIVE source data" if result.live else "DEMO data (synthetic or recorded fixtures)"
    rows: list[tuple[str, Any]] = [
        ("Profile", profile.name),
        ("Data", data_note),
        ("Screened as of (IST)", now.astimezone(IST).strftime("%d %b %Y %H:%M")),
        (
            "Sources",
            ", ".join(
                f"{r.name} ({'ok' if r.ok else 'failed'}, {r.fetched})" for r in result.sources
            ),
        ),
        ("Fetched", len(result.screened)),
        ("Duplicates removed", len(result.duplicates)),
        ("Unique", len(result.unique)),
        ("Removed by eligibility filters", len(result.ineligible)),
        ("Eligible", len(result.eligible)),
        ("High priority", len(result.tier("high"))),
        ("Medium priority", len(result.tier("medium"))),
        ("Low priority", len(result.tier("low"))),
    ]
    for k, v in rows:
        summary.append([k, v])
    for cell in summary["A"]:
        cell.font = Font(bold=True)
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 90

    def ranked(xs: list[ScreenedTender]) -> list[ScreenedTender]:
        return sorted(xs, key=lambda s: -(s.score.total if s.score else 0))

    for tier, title in (
        ("high", "High Priority"),
        ("medium", "Medium Priority"),
        ("low", "Low Priority"),
    ):
        _sheet(wb, title, BASE_COLS, [_base(s, now) for s in ranked(result.tier(tier))])

    factor_cols = [f"{f} (w={profile.weights[f]:.2f})" for f in FACTORS]
    all_rows = []
    for s in ranked(result.eligible):
        contrib = {f.name: f.contribution for f in s.score.factors} if s.score else {}
        all_rows.append([*_base(s, now), *(contrib.get(f) for f in FACTORS)])
    _sheet(wb, "All Scored", [*BASE_COLS, *factor_cols], all_rows)

    _sheet(
        wb,
        "Ineligible",
        ["Tender ID", "Portal", "Title", "Organisation", "Closing (IST)", "Reasons"],
        [
            [
                s.tender.source_tender_id,
                s.tender.portal,
                s.tender.title,
                s.tender.organisation,
                s.tender.closing_at.astimezone(IST).replace(tzinfo=None),
                "; ".join(s.eligibility.reasons) if s.eligibility else "",
            ]
            for s in result.ineligible
        ],
    )
    _sheet(
        wb,
        "Duplicates",
        ["Tender ID", "Portal", "Title", "Duplicate of", "Jaccard similarity"],
        [
            [
                s.tender.source_tender_id,
                s.tender.portal,
                s.tender.title,
                s.duplicate_of,
                s.duplicate_similarity,
            ]
            for s in result.duplicates
        ],
    )
    wb.save(path)
    return path
