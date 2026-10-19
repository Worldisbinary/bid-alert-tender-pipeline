import contextlib
import json
from datetime import timedelta

from openpyxl import load_workbook
from sqlalchemy import func, select

from bidalert.alerts import AlertDispatcher
from bidalert.pipeline import run_pipeline
from bidalert.sources.base import SourceAdapter, SourceError
from bidalert.sources.fixtures import FixtureAdapter
from bidalert.sources.synthetic import SyntheticAdapter
from bidalert.storage.schema import AlertSentRow, PipelineRunRow, TenderRow
from tests.conftest import FIXTURES, NOW


class RecordingNotifier:
    channel = "fake"

    def __init__(self):
        self.batches = []

    def send(self, alerts):
        self.batches.append(list(alerts))


class Broken(SourceAdapter):
    name = "broken"
    is_live = True
    description = "always fails"

    def fetch(self):
        raise SourceError("portal down")


def test_end_to_end_synthetic_with_reports_and_alerts(sessions, profile, tmp_path):
    notifier = RecordingNotifier()
    result = run_pipeline(
        [SyntheticAdapter(now=NOW)],
        profile,
        sessions=sessions,
        notifiers=[notifier],
        now=NOW,
        output_dir=tmp_path / "out",
        site_dir=tmp_path / "site",
    )
    st = result.stats()
    assert st["fetched"] == 180 and st["duplicates"] == 20 and st["unique"] == 160
    assert st["eligible"] + st["ineligible"] == 160 and st["eligible"] > 0
    assert st["high"] + st["medium"] + st["low"] == st["eligible"]
    assert not result.live and result.status == "success"

    wb = load_workbook(result.report_paths["excel"])
    assert wb.sheetnames == [
        "Summary",
        "High Priority",
        "Medium Priority",
        "Low Priority",
        "All Scored",
        "Ineligible",
        "Duplicates",
    ]
    assert wb["All Scored"].max_row == st["eligible"] + 1

    html = (tmp_path / "site" / "index.html").read_text("utf-8")
    assert "Demo data, not real tenders" in html and "Live public data" not in html
    summary = json.loads((tmp_path / "site" / "summary.json").read_text("utf-8"))
    assert summary["live"] is False and summary["stats"]["fetched"] == 180

    sent_first = sum(len(b) for b in notifier.batches)
    assert sent_first == result.alerts_built > 0

    # Second run, same data, 6 hours later: idempotent rows, no repeated alerts.
    again = run_pipeline(
        [SyntheticAdapter(now=NOW)],
        profile,
        sessions=sessions,
        notifiers=[notifier],
        now=NOW + timedelta(hours=6),
    )
    assert again.upsert is not None and again.upsert.inserted == 0
    with sessions() as s:
        assert s.scalar(select(func.count()).select_from(TenderRow)) == 180
        assert s.scalar(select(func.count()).select_from(PipelineRunRow)) == 2
        assert s.scalar(select(func.count()).select_from(AlertSentRow)) == sent_first
    assert sum(len(b) for b in notifier.batches) == sent_first


def test_fixtures_screen_at_capture_time(profile, tmp_path):
    result = run_pipeline([FixtureAdapter(FIXTURES / "cppp")], profile, site_dir=tmp_path)
    assert result.now == FixtureAdapter(FIXTURES / "cppp").recorded_at
    assert len(result.screened) == 20
    assert result.reason_counts()["excluded keyword"] >= 1  # e.g. the canteen civil work tender
    assert all(s.eligibility and "tender value" in s.eligibility.unverified for s in result.unique)


def test_failed_source_is_reported_not_fatal(sessions, profile):
    result = run_pipeline(
        [Broken(), SyntheticAdapter(now=NOW)], profile, sessions=sessions, now=NOW
    )
    assert result.status == "partial"
    assert result.sources[0].error and "portal down" in result.sources[0].error
    assert not result.live
    with sessions() as s:
        assert s.scalars(select(PipelineRunRow.status)).one() == "partial"


def test_live_flag_needs_a_live_source_with_data(profile, tmp_path):
    class FakeLive(SyntheticAdapter):
        name = "fake-live"
        is_live = True

    result = run_pipeline([FakeLive(now=NOW, n=10)], profile, now=NOW, site_dir=tmp_path)
    assert result.live
    assert "Live public data" in (tmp_path / "index.html").read_text("utf-8")


def test_crash_is_recorded_as_failed_run(sessions, profile, monkeypatch):
    def boom(*a, **k):
        raise RuntimeError("disk full")

    monkeypatch.setattr("bidalert.pipeline.TenderRepository.upsert", boom)
    with contextlib.suppress(RuntimeError):
        run_pipeline([SyntheticAdapter(now=NOW, n=5)], profile, sessions=sessions, now=NOW)
    with sessions() as s:
        run = s.scalars(select(PipelineRunRow)).one()
        assert run.status == "failed" and "disk full" in (run.error or "")


def test_dispatcher_isolates_broken_channel():
    class Ledger:
        def __init__(self):
            self.rows = set()

        def already_sent(self, t, c, k):
            return (t, c, k) in self.rows

        def record(self, t, c, k, at):
            self.rows.add((t, c, k))

    class Failing:
        channel = "bad"

        def send(self, alerts):
            raise RuntimeError("smtp down")

    from bidalert.alerts import Alert

    good = RecordingNotifier()
    alert = Alert(1, "new_high_priority", "t", "o", NOW, 0.8, None)
    res = AlertDispatcher([Failing(), good], Ledger()).dispatch([alert], NOW)
    assert res.failed_channels == ("bad",) and res.sent == {"fake": 1}
